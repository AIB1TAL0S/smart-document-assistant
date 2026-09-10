import io
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from engine import Engine, extract


class IndexTests(unittest.TestCase):
    def test_index_survives_restart_and_delete_removes_evidence(self):
        with tempfile.TemporaryDirectory() as folder:
            engine = Engine(Path(folder))
            with patch.object(engine, 'embeddings', return_value=[[1.0, 0.0]]):
                doc = engine.ingest('policy.txt', b'Refunds are available for 30 days.')
            recovered = Engine(Path(folder))
            self.assertEqual(recovered.documents(), [doc])
            with patch.object(recovered, 'embeddings', return_value=[[1.0, 0.0]]), patch.object(
                recovered, 'request', return_value={'message': {'content': '30 days [1].'}}
            ):
                answer = recovered.ask('How long are refunds available?')
            self.assertEqual(answer['sources'][0]['text'], 'Refunds are available for 30 days.')
            self.assertTrue(recovered.delete(doc['id']))
            with self.assertRaisesRegex(ValueError, 'Upload a document'):
                recovered.ask('How long are refunds available?')

    def test_embedding_failure_does_not_leave_partial_document(self):
        with tempfile.TemporaryDirectory() as folder:
            engine = Engine(Path(folder))
            with patch.object(engine, 'embeddings', side_effect=RuntimeError('Model unavailable')):
                with self.assertRaises(RuntimeError):
                    engine.ingest('notes.txt', b'Important document text')
            self.assertEqual(Engine(Path(folder)).documents(), [])

    def test_embedding_model_change_requires_empty_index(self):
        with tempfile.TemporaryDirectory() as folder:
            engine = Engine(Path(folder))
            with patch.object(engine, 'embeddings', return_value=[[1.0, 0.0]]):
                doc = engine.ingest('notes.txt', b'First model evidence.')
            with patch.dict(os.environ, {'EMBED_MODEL': 'another-local-model'}):
                changed = Engine(Path(folder))
            with self.assertRaisesRegex(RuntimeError, 'different embedding model'):
                changed.ask('What evidence?')
            changed.delete(doc['id'])
            with patch.object(changed, 'embeddings', return_value=[[0.0, 1.0]]):
                changed.ingest('new.txt', b'New model evidence.')

    def test_docx_tables_are_searchable(self):
        doc = Document()
        doc.add_paragraph('Travel policy')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Meal allowance'
        table.cell(0, 1).text = '$42 per day'
        stream = io.BytesIO()
        doc.save(stream)
        text = ' '.join(chunk for _, chunk in extract('policy.docx', stream.getvalue()))
        self.assertIn('Meal allowance', text)
        self.assertIn('$42 per day', text)

    def test_empty_or_invalid_documents_are_rejected(self):
        for name, content in [('empty.txt', b' \n '), ('bad.pdf', b'not a PDF'), ('bad.docx', b'broken zip'), ('script.html', b'<p>hello</p>')]:
            with self.subTest(name=name), self.assertRaises(ValueError):
                extract(name, content)


if __name__ == '__main__':
    unittest.main()
